import io
import json
import unittest
from contextlib import closing
from datetime import datetime
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
import app
from tests import test_claim_correctness as existing


class SplitImportTests(unittest.TestCase):
    def setUp(self):
        existing.ClaimCorrectnessTests.setUp(self)
        # Excel 类型判断需要真实 datetime 类型，而非旧测试的 MagicMock。
        clock = patch.object(app, 'datetime', datetime)
        clock.start()
        self.addCleanup(clock.stop)
    request = existing.ClaimCorrectnessTests.request
    payment = existing.ClaimCorrectnessTests.payment
    active = existing.ClaimCorrectnessTests.active

    def parse(self, text):
        return app.parse_split_import(text.encode('utf-8-sig'), '分摊.csv')

    def upload(self, pid, content, authenticated=True, filename='split.csv'):
        boundary = 'split-import-test-boundary'
        body = (f'--{boundary}\r\nContent-Disposition: form-data; name="attachment"; filename="{filename}"\r\n'
                'Content-Type: application/octet-stream\r\n\r\n').encode() + content + f'\r\n--{boundary}--\r\n'.encode()
        return self.request('POST', f'/split-claim/{pid}/import-preview', authenticated=authenticated,
                            raw_body=body, content_type=f'multipart/form-data; boundary={boundary}',
                            extra_headers=[(b'accept', b'application/json')])

    def test_exact_matching_constraints_and_ambiguous_names(self):
        with patch.object(app, 'CATALOG', {'部门甲': {'中心甲': ['唯一产品', '同名产品']}, '部门乙': {'中心乙': ['同名产品']}}):
            result = self.parse('部门,中心/小组,产品名称,金额（元）,备注\n,,唯一产品,100,收入\n,,同名产品,-10,退款\n部门乙,中心乙,同名产品,20,补款\n部门乙,,唯一产品,5,冲突\n,,未知产品,10,未知\n,,合计,125,\n')
        self.assertEqual(result['total_cents'], 12500)
        self.assertEqual(result['unmatched'], 3)
        self.assertEqual(result['skipped_totals'], 1)
        self.assertEqual(result['lines'][0]['department'], '部门甲')
        self.assertEqual(result['lines'][1]['project'], '')
        self.assertEqual(result['lines'][1]['amount_cents'], -1000)
        self.assertEqual(result['lines'][2]['department'], '部门乙')
        self.assertEqual(result['lines'][3]['project'], '')

    def test_reject_bad_amounts_ambiguous_headers_and_missing_rows(self):
        for text in ['产品,金额\nA,1.001', '产品,金额\nA,0', '产品,金额\nA,NaN',
                     '产品,金额\nA,=1+2', '产品,金额\n,10', '产品,金额\nA,',
                     '产品,金额,分摊金额\nA,100,80', '商品名称,实付金额,退款金额\nA,100,10',
                     '产品,金额\nA,1\n,', '产品,金额\nA,1\n未知列说明,']:
            if text.endswith('\n,'):
                self.assertEqual(len(self.parse(text)['lines']), 1)
                continue
            with self.subTest(text=text), self.assertRaises(app.HTTPException):
                self.parse(text)

    def test_xlsx_all_sheets_formula_and_invalid_sheet(self):
        book = Workbook()
        sheet = book.active
        sheet.append(['月末分摊'])
        sheet.append(['项目名称', '分摊金额'])
        sheet.append([self.project, 123.45])
        other = book.create_sheet('退款')
        other.append(['产品', '金额'])
        other.append([self.project, -23.45])
        out = io.BytesIO()
        book.save(out)
        result = app.parse_split_import(out.getvalue(), 'test.xlsx')
        self.assertEqual(result['total_cents'], 10000)
        self.assertEqual(len(result['lines']), 2)
        other['B2'] = '=1+2'
        out = io.BytesIO(); book.save(out)
        with self.assertRaisesRegex(app.HTTPException, '公式'):
            app.parse_split_import(out.getvalue(), 'test.xlsx')
        other['B2'] = 10
        other['A1'] = '不支持的标题'
        out = io.BytesIO(); book.save(out)
        with self.assertRaisesRegex(app.HTTPException, '未找到表头'):
            app.parse_split_import(out.getvalue(), 'test.xlsx')
        book.close()

    def test_encoding_limits_and_corrupt_files(self):
        result = app.parse_split_import('产品,金额\n教育,100'.encode('gb18030'), 'test.csv')
        self.assertEqual(result['lines'][0]['source_project'], '教育')
        for data, name in [(b'garbage', 'test.xlsx'), (b'abc', 'test.pdf'), (b'', 'test.csv'),
                           (b'x' * (app.SPLIT_IMPORT_MAX_BYTES + 1), 'test.csv'),
                           (('产品,金额\n' + 'A,1\n' * 2001).encode(), 'test.csv')]:
            with self.subTest(name=name, size=len(data)), self.assertRaises(app.HTTPException):
                app.parse_split_import(data, name)

    def test_preview_is_read_only_then_existing_submit_and_stale_balance(self):
        pid = self.payment(amount=20000, claims=(5000,))
        raw = f'部门,中心/小组,产品,金额,备注\n{self.department},{self.team},{self.project},160,收入\n{self.department},{self.team},{self.project},-10,退款'.encode()
        response = self.upload(pid, raw)
        self.assertEqual(response.status_code, 200, response.text)
        result = json.loads(response.text)
        self.assertEqual(result['remaining_cents'], 15000)
        self.assertEqual(result['total_cents'], 15000)
        self.assertEqual(self.active(pid), 5000)
        with closing(app.get_conn()) as conn:
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM notification_outbox').fetchone()[0], 0)
        # 预览后另一人认领了一部分，提交仍按数据库实时余额拒绝。
        data = dict(departments=[self.department], teams=[self.team], projects=[self.project], amounts=['10'], notes=[''])
        self.assertEqual(self.request('POST', f'/split-claim/{pid}', data).status_code, 303)
        lines = result['lines']
        submit = dict(departments=[r['department'] for r in lines], teams=[r['team'] for r in lines],
                      projects=[r['project'] for r in lines], amounts=[r['amount'] for r in lines], notes=[r['note'] for r in lines])
        self.assertEqual(self.request('POST', f'/split-claim/{pid}', submit).status_code, 409)
        self.assertEqual(self.active(pid), 6000)
        new_pid = self.payment(amount=15000)
        self.assertEqual(self.request('POST', f'/split-claim/{new_pid}', submit).status_code, 303)
        self.assertEqual(self.active(new_pid), 15000)

    def test_two_thousand_lines_submit_through_existing_form(self):
        row = f'{self.department},{self.team},{self.project},0.01'
        result = self.parse('部门,中心/小组,产品,金额\n' + '\n'.join([row] * 2000))
        self.assertEqual(result['total_cents'], 2000)
        pid = self.payment(amount=2000)
        lines = result['lines']
        data = dict(departments=[r['department'] for r in lines], teams=[r['team'] for r in lines],
                    projects=[r['project'] for r in lines], amounts=[r['amount'] for r in lines], notes=[''] * 2000)
        self.assertEqual(self.request('POST', f'/split-claim/{pid}', data).status_code, 303)
        self.assertEqual(self.active(pid), 2000)

    def test_auth_payment_status_and_malformed_file(self):
        pid = self.payment()
        raw = b'product,amount\nA,1'
        self.assertEqual(self.upload(pid, raw, authenticated=False).status_code, 401)
        self.assertEqual(self.upload(99999, raw).status_code, 404)
        closed = self.payment(status='closed')
        self.assertEqual(self.upload(closed, raw).status_code, 409)
        self.assertEqual(self.upload(pid, raw).status_code, 400)
        self.assertEqual(self.active(pid), 0)

    def test_import_ui_and_template(self):
        pid = self.payment(amount=20000, claims=(5000,))
        response = self.request('GET', f'/split-claim?q=%23{pid}')
        self.assertEqual(response.status_code, 200)
        self.assertIn('识别并填入明细', response.text)
        self.assertIn('data-remaining-cents="15000"', response.text)
        request = app.Request({'type':'http', 'headers':[]})
        with patch.object(app, 'actor_from_request', return_value=self.actor):
            template = app.split_import_template(request)
        workbook = load_workbook(io.BytesIO(template.body))
        self.assertEqual(list(workbook.active.values), [('部门', '中心/小组', '产品名称', '分摊金额', '备注')])
        workbook.close()


if __name__ == '__main__':
    unittest.main()
